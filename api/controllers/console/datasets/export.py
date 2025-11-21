import csv
import io
import json
from urllib.parse import quote

from flask import Response
from flask_login import current_user
from flask_restful import Resource, reqparse
from werkzeug.exceptions import BadRequest, Forbidden, NotFound

from controllers.console import api
from controllers.console.wraps import (
    account_initialization_required,
    setup_required,
)
from extensions.ext_database import db
from libs.login import login_required
from models.dataset import Document, DocumentSegment
from services.dataset_service import DatasetService


def _safe_filename(filename: str) -> str:
    try:
        filename.encode("ascii")
        return f"filename=\"{filename}\""
    except UnicodeEncodeError:
        encoded = quote(filename, safe="")
        safe_ascii = "".join(c if ord(c) < 128 else "_" for c in filename)
        return f"filename=\"{safe_ascii}\"; filename*=UTF-8''{encoded}"


class ConsoleDatasetExportApi(Resource):
    @setup_required
    @login_required
    @account_initialization_required
    def get(self, dataset_id):
        dataset_id = str(dataset_id)

        dataset = DatasetService.get_dataset(dataset_id)
        if dataset is None:
            raise NotFound("Dataset not found.")
        try:
            DatasetService.check_dataset_permission(dataset, current_user)
        except Exception as e:
            raise Forbidden(str(e))

        # 参数
        parser = reqparse.RequestParser()
        parser.add_argument(
            "format",
            type=str,
            default="json",
            choices=["json", "csv", "txt", "docx", "word"],
            location="args",
        )
        parser.add_argument("document_ids", type=str, action="append", location="args")
        args = parser.parse_args()

        export_format = args["format"]
        document_ids = args["document_ids"]

        # 查询文档
        query = db.session.query(Document).filter_by(dataset_id=dataset_id, tenant_id=current_user.current_tenant_id)
        if document_ids:
            query = query.filter(Document.id.in_(document_ids))
        documents = query.order_by(Document.created_at.desc()).all()
        if not documents:
            raise NotFound("No documents found in the dataset.")

        if export_format == "json":
            return self._export_json(dataset, documents)
        elif export_format == "csv":
            return self._export_csv(dataset, documents)
        elif export_format == "txt":
            return self._export_txt(dataset, documents)
        elif export_format in ("docx", "word"):
            return self._export_docx(dataset, documents)
        else:
            raise BadRequest("Unsupported export format.")

    @staticmethod
    def _export_json(dataset, documents):
        export_data = {
            "dataset": {
                "id": dataset.id,
                "name": dataset.name,
                "description": dataset.description,
                "created_at": dataset.created_at.isoformat(),
                "updated_at": dataset.updated_at.isoformat(),
            },
            "documents": [],
        }

        for document in documents:
            doc_data = {
                "id": document.id,
                "name": document.name,
                "data_source_type": document.data_source_type,
                "indexing_status": document.indexing_status,
                "created_at": document.created_at.isoformat(),
                "updated_at": document.updated_at.isoformat(),
                "enabled": document.enabled,
                "segment_count": document.segment_count,
                "doc_type": document.doc_type,
                "doc_metadata": getattr(document, "doc_metadata_details", None),
            }

            segments = (
                db.session.query(DocumentSegment)
                .filter_by(document_id=document.id, tenant_id=document.tenant_id)
                .order_by(DocumentSegment.position)
                .all()
            )

            doc_data["segments"] = []
            for segment in segments:
                segment_data = {
                    "id": segment.id,
                    "position": segment.position,
                    "content": segment.content,
                    "answer": segment.answer,
                    "document_name": document.name,
                    "word_count": segment.word_count,
                    "status": segment.status,
                    "keywords": segment.keywords,
                    "created_at": segment.created_at.isoformat(),
                    "updated_at": segment.updated_at.isoformat(),
                }
                doc_data["segments"].append(segment_data)

            export_data["documents"].append(doc_data)

        filename = f"dataset_{dataset.name}_{dataset.id}.json"
        response = Response(
            json.dumps(export_data, ensure_ascii=False, indent=2),
            mimetype="application/json",
            headers={
                "Content-Disposition": f"attachment; {_safe_filename(filename)}",
                "Content-Type": "application/json; charset=utf-8",
            },
        )
        return response

    @staticmethod
    def _export_csv(dataset, documents):
        output = io.StringIO()
        writer = csv.writer(output)

        writer.writerow(["Dataset Information"])
        writer.writerow(["ID", dataset.id])
        writer.writerow(["Name", dataset.name])
        writer.writerow(["Description", dataset.description])
        writer.writerow(["Created At", dataset.created_at.isoformat()])
        writer.writerow([])

        writer.writerow(["Documents"])
        writer.writerow([
            "Document ID",
            "Name",
            "Data Source Type",
            "Indexing Status",
            "Created At",
            "Updated At",
            "Enabled",
            "Segment Count",
            "Doc Type",
        ])

        for document in documents:
            writer.writerow(
                [
                    document.id,
                    document.name,
                    document.data_source_type,
                    document.indexing_status,
                    document.created_at.isoformat(),
                    document.updated_at.isoformat(),
                    document.enabled,
                    document.segment_count,
                    document.doc_type,
                ]
            )

        writer.writerow([])
        writer.writerow(["Segments"])
        # 固定前四列：ID、Position、Content、Answer，后续附加其它字段
        writer.writerow([
            "Segment ID",
            "Position",
            "Content",
            "Answer",
            "Document ID",
            "Document Name",
            "Word Count",
            "Status",
            "Keywords",
            "Created At",
            "Updated At",
        ])

        for document in documents:
            segments = (
                db.session.query(DocumentSegment)
                .filter_by(document_id=document.id, tenant_id=document.tenant_id)
                .order_by(DocumentSegment.position)
                .all()
            )
            for segment in segments:
                keywords_str = json.dumps(segment.keywords, ensure_ascii=False) if segment.keywords else ""
                writer.writerow(
                    [
                        segment.id,
                        segment.position,
                        segment.content,
                        (segment.answer or ""),
                        document.id,
                        document.name,
                        segment.word_count,
                        segment.status,
                        keywords_str,
                        segment.created_at.isoformat(),
                        segment.updated_at.isoformat(),
                    ]
                )

        filename = f"dataset_{dataset.name}_{dataset.id}.csv"
        response = Response(
            output.getvalue(),
            mimetype="text/csv",
            headers={"Content-Disposition": f"attachment; {_safe_filename(filename)}"},
        )
        return response

    @staticmethod
    def _export_txt(dataset, documents):
        output = io.StringIO()
        output.write(f"Dataset: {dataset.name}\n")
        output.write(f"ID: {dataset.id}\n")
        output.write(f"Description: {dataset.description}\n")
        output.write(f"Created At: {dataset.created_at.isoformat()}\n")
        output.write(f"Updated At: {dataset.updated_at.isoformat()}\n")
        output.write("=" * 80 + "\n\n")

        for document in documents:
            output.write(f"Document: {document.name}\n")
            output.write(f"ID: {document.id}\n")
            output.write(f"Data Source Type: {document.data_source_type}\n")
            output.write(f"Indexing Status: {document.indexing_status}\n")
            output.write(f"Created At: {document.created_at.isoformat()}\n")
            output.write(f"Updated At: {document.updated_at.isoformat()}\n")
            output.write(f"Enabled: {document.enabled}\n")
            output.write(f"Segment Count: {document.segment_count}\n")
            if document.doc_type:
                output.write(f"Doc Type: {document.doc_type}\n")
            output.write("-" * 40 + "\n")

            segments = (
                db.session.query(DocumentSegment)
                .filter_by(document_id=document.id, tenant_id=document.tenant_id)
                .order_by(DocumentSegment.position)
                .all()
            )
            for i, segment in enumerate(segments, 1):
                output.write(f"Segment {i}\n")
                output.write(f"ID: {segment.id}\n")
                output.write(f"Position: {segment.position}\n")
                output.write(f"Document Name: {document.name}\n")
                output.write("Content:\n")
                output.write(segment.content)
                output.write("\n")
                output.write(f"Answer: {segment.answer or ''}\n")
                output.write(f"Word Count: {segment.word_count}\n")
                output.write(f"Status: {segment.status}\n")
                if segment.keywords:
                    output.write(f"Keywords: {json.dumps(segment.keywords, ensure_ascii=False)}\n")
                output.write(f"Created At: {segment.created_at.isoformat()}\n")
                output.write(f"Updated At: {segment.updated_at.isoformat()}\n")
                output.write("-" * 40 + "\n\n")

            output.write("\n" + "=" * 80 + "\n\n")

        filename = f"dataset_{dataset.name}_{dataset.id}.txt"
        response = Response(
            output.getvalue(),
            mimetype="text/plain",
            headers={"Content-Disposition": f"attachment; {_safe_filename(filename)}"},
        )
        return response

    @staticmethod
    def _export_docx(dataset, documents):
        try:
            from docx import Document as DocxDocument
        except Exception:
            raise BadRequest("python-docx 未安装，请先在后端环境安装：pip install python-docx")

        buffer = io.BytesIO()
        doc = DocxDocument()

        first_segment = True

        for d in documents:
            segments = (
                db.session.query(DocumentSegment)
                .filter_by(document_id=d.id, tenant_id=d.tenant_id)
                .order_by(DocumentSegment.position)
                .all()
            )
            for seg in segments:
                if not first_segment:
                    doc.add_paragraph("")
                first_segment = False
                doc.add_paragraph(f"Q: {seg.content or ''}")
                doc.add_paragraph(f"A: {seg.answer or ''}")

        doc.save(buffer)
        buffer.seek(0)

        filename = f"dataset_{dataset.name}_{dataset.id}.docx"
        return Response(
            buffer.getvalue(),
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f"attachment; {_safe_filename(filename)}"},
        )


api.add_resource(ConsoleDatasetExportApi, "/datasets/<uuid:dataset_id>/export")
